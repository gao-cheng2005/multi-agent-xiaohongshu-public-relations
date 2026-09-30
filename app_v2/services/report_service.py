"""报表生成。

把当前的统计数据整理成一份可下载的报告，支持两种格式：
Markdown（文字+表格，适合阅读）和 Excel（多个工作表，适合进一步分析）。

报告内容包含：数据总览、回复效果、热词、中高风险笔记、最近采集笔记。
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from sqlalchemy.ext.asyncio import AsyncSession

from .alert_service import AlertService
from .analytics_service import AnalyticsService
from .content_service import ContentService


def _now_text() -> str:
    """生成当前时间的展示文字，用于报告开头。"""
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _filename(ext: str) -> str:
    """生成带时间戳的文件名，避免重名。"""
    return f"舆情分析报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"


class ReportService:
    """报表生成服务。"""

    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def _snapshot(self) -> dict:
        """一次性收集生成报告所需的全部数据。

        复用内容、统计、预警三个服务来取数据，避免到处拼 SQL。
        """
        content_service = ContentService(self.session)
        analytics_service = AnalyticsService(self.session)
        alert_service = AlertService(self.session)
        return {
            "generated_at": _now_text(),
            "collect_mode": "opencli",
            "stats": await content_service.stats(),
            "effect": await analytics_service.effect(),
            "hotwords": await analytics_service.hotwords(20),
            "alerts": await alert_service.list_alerts(),
            "contents": await content_service.list_contents(50),
        }

    async def build_markdown(self) -> tuple[str, str, str]:
        """生成 Markdown 报告，返回（内容, 类型, 文件名）。"""
        snap = await self._snapshot()
        s = snap["stats"]
        e = snap["effect"]
        sentiment = s.get("sentiment", {})
        risk = s.get("risk", {})

        # 逐行拼装报告文字
        lines = [
            "# 舆情公关分析报告",
            "",
            f"- 生成时间：{snap['generated_at']}",
            "- 采集模式：真实抓取(opencli)",
            "",
            "## 一、数据总览",
            "",
            "| 指标 | 数值 |",
            "|------|------|",
            f"| 笔记数 | {s.get('notes', 0)} |",
            f"| 评论数 | {s.get('comments', 0)} |",
            f"| 待处理预警 | {s.get('open_alerts', 0)} |",
            f"| 高风险预警 | {s.get('high_alerts', 0)} |",
            f"| 正面 / 中性 / 负面 | {sentiment.get('正面', 0)} / {sentiment.get('中性', 0)} / {sentiment.get('负面', 0)} |",
            f"| 低 / 中 / 高风险 | {risk.get('低', 0)} / {risk.get('中', 0)} / {risk.get('高', 0)} |",
            "",
            "## 二、回复效果",
            "",
            "| 指标 | 数值 |",
            "|------|------|",
            f"| 应对方案总数 | {e.get('plans_total', 0)} |",
            f"| 待审核 | {e.get('plans_pending', 0)} |",
            f"| 已采纳 | {e.get('plans_accepted', 0)} |",
            f"| 已关闭 | {e.get('plans_closed', 0)} |",
            f"| 采纳率 | {round(e.get('adoption_rate', 0) * 100, 1)}% |",
            f"| 平均决策耗时 | {round(e.get('avg_decision_seconds', 0), 1)} 秒 |",
            f"| 预警处理率 | {round(e.get('alert_handle_rate', 0) * 100, 1)}% |",
            "",
            "## 三、热词 Top20",
            "",
        ]

        # 热词表格
        hotwords = snap["hotwords"]
        if hotwords:
            lines.extend(["| 排名 | 热词 | 出现次数 |", "|------|------|----------|"])
            for i, hw in enumerate(hotwords, 1):
                lines.append(f"| {i} | {hw['word']} | {hw['count']} |")
        else:
            lines.append("暂无热词数据。")

        # 中高风险笔记表格
        lines.extend(["", "## 四、中高风险笔记", ""])
        if snap["alerts"]:
            lines.extend(["| 笔记 | 风险 | 状态 |", "|------|------|------|"])
            for a in snap["alerts"][:30]:
                status = "待处理" if a.get("open_count", 0) > 0 else "已处理"
                lines.append(f"| {a.get('note_title', '') or '—'} | {a.get('level', '低')} | {status} |")
        else:
            lines.append("暂无中高风险笔记。")

        # 最近采集笔记表格
        lines.extend(["", "## 五、最近采集笔记", ""])
        if snap["contents"]:
            lines.extend(["| 标题 | 情感 | 风险 | 点赞 | 评论 |", "|------|------|------|------|------|"])
            for c in snap["contents"][:30]:
                lines.append(
                    f"| {c.get('title', '') or '—'} | {c.get('sentiment', '') or '—'} | "
                    f"{c.get('risk_level', '') or '—'} | {c.get('like_count', 0)} | {c.get('comment_count', 0)} |"
                )
        else:
            lines.append("暂无采集笔记。")

        return "\n".join(lines), "text/markdown; charset=utf-8", _filename("md")

    async def build_xlsx(self) -> tuple[bytes, str, str]:
        """生成 Excel 报告，返回（内容, 类型, 文件名）。"""
        from openpyxl import Workbook
        from openpyxl.styles import Font

        snap = await self._snapshot()
        s = snap["stats"]
        e = snap["effect"]

        # 创建工作簿和第一个工作表"总览"
        wb = Workbook()
        ws = wb.active
        ws.title = "总览"
        ws.append(["指标", "数值"])
        for label, value in [
            ("生成时间", snap["generated_at"]),
            ("采集模式", "真实抓取(opencli)"),
            ("笔记数", s.get("notes", 0)),
            ("评论数", s.get("comments", 0)),
            ("待处理预警", s.get("open_alerts", 0)),
            ("高风险预警", s.get("high_alerts", 0)),
        ]:
            ws.append([label, value])

        # 工作表"回复效果"
        ws2 = wb.create_sheet("回复效果")
        ws2.append(["指标", "数值"])
        for label, value in [
            ("应对方案总数", e.get("plans_total", 0)),
            ("待审核", e.get("plans_pending", 0)),
            ("已采纳", e.get("plans_accepted", 0)),
            ("已关闭", e.get("plans_closed", 0)),
            ("采纳率", round(e.get("adoption_rate", 0) * 100, 1)),
            ("平均决策耗时(秒)", round(e.get("avg_decision_seconds", 0), 1)),
            ("预警处理率(%)", round(e.get("alert_handle_rate", 0) * 100, 1)),
        ]:
            ws2.append([label, value])

        # 工作表"热词"
        ws3 = wb.create_sheet("热词")
        ws3.append(["排名", "热词", "出现次数"])
        for i, hw in enumerate(snap["hotwords"], 1):
            ws3.append([i, hw["word"], hw["count"]])

        # 工作表"中高风险笔记"
        ws4 = wb.create_sheet("中高风险笔记")
        ws4.append(["笔记", "风险", "状态"])
        for a in snap["alerts"][:100]:
            ws4.append([
                a.get("note_title", "") or "—",
                a.get("level", "低"),
                "待处理" if a.get("open_count", 0) > 0 else "已处理",
            ])

        # 工作表"最近采集笔记"
        ws5 = wb.create_sheet("最近采集笔记")
        ws5.append(["标题", "情感", "风险", "点赞", "评论"])
        for c in snap["contents"][:100]:
            ws5.append([
                c.get("title", "") or "—",
                c.get("sentiment", "") or "—",
                c.get("risk_level", "") or "—",
                c.get("like_count", 0),
                c.get("comment_count", 0),
            ])

        # 给每个表的第一行加粗，更美观
        for sheet in (ws, ws2, ws3, ws4, ws5):
            for cell in sheet[1]:
                cell.font = Font(bold=True)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _filename("xlsx")
